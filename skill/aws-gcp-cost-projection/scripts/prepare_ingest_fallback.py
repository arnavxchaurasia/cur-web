#!/usr/bin/env python3
"""
prepare_ingest_fallback.py

Runs ONLY when the deterministic ingest.py parsers all failed to recognize
the raw bill file (Phase 1 retry path). Zero LLM tokens — this is the
"auditor" half of the auditor/LLM split: it dumps a bounded, readable preview
of the raw file (every sheet if Excel, every line up to a cap if CSV/text) to
projection-audit/ingest_fallback_manifest.md so the LLM only has to read one
file and never has to guess at parsing the raw bytes itself.

The LLM's job (see the Phase 1 fallback prompt in orchestrate.go) is to read
that manifest and write projection-audit/normalized_bill.csv in the flat-CSV
schema ingest.py already understands (Description, Region, Usage Quantity,
Amount in USD) — it never touches the database directly.
"""
import os
import sys
import glob

JOB_DIR = os.getcwd()
OUT_DIR = os.path.join(JOB_DIR, "projection-audit")
OUT_PATH = os.path.join(OUT_DIR, "ingest_fallback_manifest.md")

MAX_ROWS_PER_SHEET = 500
MAX_TEXT_LINES = 800


def _preview_excel(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = [f"# Excel workbook: {os.path.basename(path)}", ""]
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"## Sheet: {sheet_name}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS_PER_SHEET:
                lines.append(f"... ({ws.max_row - MAX_ROWS_PER_SHEET} more rows truncated)")
                break
            if all(v is None for v in row):
                continue
            lines.append(",".join("" if v is None else str(v) for v in row))
        lines.append("")
    return "\n".join(lines)


def _preview_text(path):
    with open(path, "r", errors="replace") as f:
        text_lines = f.readlines()
    lines = [f"# Text/CSV file: {os.path.basename(path)}", ""]
    for i, line in enumerate(text_lines):
        if i >= MAX_TEXT_LINES:
            lines.append(f"... ({len(text_lines) - MAX_TEXT_LINES} more lines truncated)")
            break
        lines.append(line.rstrip("\n"))
    return "\n".join(lines)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    inputs = glob.glob(os.path.join(JOB_DIR, "input.*"))
    if not inputs:
        sys.stderr.write("prepare_ingest_fallback: no input.* file found — nothing to preview\n")
        sys.exit(0)
    input_file = inputs[0]
    ext = os.path.splitext(input_file)[1].lower()
    try:
        if ext in (".xlsx", ".xls"):
            content = _preview_excel(input_file)
        else:
            content = _preview_text(input_file)
    except Exception as e:
        content = f"# Could not preview {os.path.basename(input_file)}: {e}\n"

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"prepare_ingest_fallback: wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
